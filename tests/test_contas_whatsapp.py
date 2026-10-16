import asyncio
from datetime import date, datetime

import pytest

from finbot import contas_whatsapp, ia, whatsapp
from finbot.config import Config
from finbot.db import Database
from finbot.whatsapp import ColetaWhatsApp, MensagemWhatsApp

HOJE = date(2026, 10, 16)
AGORA = datetime(2026, 10, 16, 8, 0)
PNG = b"\x89PNG\r\n\x1a\n"


@pytest.fixture(autouse=True)
def sem_esperas(monkeypatch):
    monkeypatch.setattr(contas_whatsapp, "PAUSA_ENTRE_CHAMADAS_S", 0)
    monkeypatch.setattr(contas_whatsapp, "ESPERA_LIMITE_S", 0)


@pytest.fixture
def cfg(tmp_path):
    return Config(telegram_token="x", whatsapp_grupo="Casa", whatsapp_perfil=tmp_path / "perfil")


@pytest.fixture
def db(tmp_path):
    return Database(tmp_path / "teste.db")


def _msg(chave, texto="", remetente="Maria Eduarda", print_png=None):
    return MensagemWhatsApp(
        chave=chave,
        momento=datetime(2026, 10, 16, 9, 15),
        remetente=remetente,
        texto=texto,
        print_png=print_png,
    )


def _coleta_falsa(monkeypatch, coleta: ColetaWhatsApp) -> None:
    async def fake(*args, **kwargs):
        return coleta

    monkeypatch.setattr(whatsapp, "coletar_grupo", fake)


async def _sem_qr(png: bytes) -> None:
    pass


def _rodar(cfg, db):
    return asyncio.run(contas_whatsapp.sincronizar(cfg, db, AGORA, HOJE, _sem_qr))


def _contas_out(db):
    return db.listar_contas_a_pagar(date(2026, 10, 1), date(2026, 10, 31))


def test_texto_cria_contas_e_nao_reprocessa_as_mensagens(monkeypatch, cfg, db):
    _coleta_falsa(
        monkeypatch,
        ColetaWhatsApp("ok", [_msg("a", "Luz vence dia 15, R$ 180"), _msg("b", "Internet paga", "Gabriel")]),
    )
    chamadas = []

    async def fake_texto(linhas, hoje):
        chamadas.append(linhas)
        return [
            ia.ContaExtraida(descricao="Luz", valor=180.0, vencimento="2026-10-15"),
            ia.ContaExtraida(descricao="Internet", valor=99.9, vencimento="2026-10-20", pago=True),
        ]

    monkeypatch.setattr(ia, "extrair_contas_de_texto", fake_texto)

    resultado = _rodar(cfg, db)
    assert resultado.status == "ok"
    assert "2 nova(s)" in resultado.texto
    assert len(chamadas) == 1 and len(chamadas[0]) == 2
    assert [(c.descricao, c.valor_centavos, c.pago) for c in _contas_out(db)] == [
        ("Luz", 18000, False),
        ("Internet", 9990, True),
    ]

    # Segunda leitura com as mesmas mensagens: a IA não é chamada de novo.
    _rodar(cfg, db)
    assert len(chamadas) == 1


def test_precisa_login_nao_chama_a_ia_e_avisa(monkeypatch, cfg, db):
    _coleta_falsa(monkeypatch, ColetaWhatsApp("precisa_login"))

    async def nao_deve_chamar(*args, **kwargs):
        raise AssertionError("a IA não deveria ser chamada sem leitura")

    monkeypatch.setattr(ia, "extrair_contas_de_texto", nao_deve_chamar)

    resultado = _rodar(cfg, db)
    assert resultado.status == "precisa_login"
    assert "QR code" in resultado.texto


def test_falha_da_ia_nao_marca_mensagem_como_lida(monkeypatch, cfg, db):
    _coleta_falsa(monkeypatch, ColetaWhatsApp("ok", [_msg("a", "Luz vence dia 15")]))

    async def falha(linhas, hoje):
        raise RuntimeError("erro qualquer")

    monkeypatch.setattr(ia, "extrair_contas_de_texto", falha)
    resultado = _rodar(cfg, db)
    assert "falharam" in resultado.texto
    assert not db.mensagem_processada("a")

    async def ok(linhas, hoje):
        return [ia.ContaExtraida(descricao="Luz", valor=180, vencimento="2026-10-15")]

    monkeypatch.setattr(ia, "extrair_contas_de_texto", ok)
    _rodar(cfg, db)
    assert db.mensagem_processada("a")
    assert len(_contas_out(db)) == 1


def test_limite_de_429_da_ia_e_tentado_de_novo(monkeypatch, cfg, db):
    _coleta_falsa(monkeypatch, ColetaWhatsApp("ok", [_msg("a", "Luz vence dia 15")]))
    tentativas = []

    async def limitado(linhas, hoje):
        tentativas.append(1)
        if len(tentativas) == 1:
            raise RuntimeError("Client error '429 Too Many Requests'")
        return [ia.ContaExtraida(descricao="Luz", valor=180, vencimento="2026-10-15")]

    monkeypatch.setattr(ia, "extrair_contas_de_texto", limitado)
    _rodar(cfg, db)
    assert len(tentativas) == 2
    assert db.mensagem_processada("a")


def test_print_so_sai_quando_a_imagem_tem_conta(monkeypatch, cfg, db):
    _coleta_falsa(
        monkeypatch,
        ColetaWhatsApp(
            "ok",
            [
                _msg("boleto", print_png=b"boleto"),
                _msg("foto", remetente="Gabriel", print_png=b"foto"),
            ],
        ),
    )
    recebidas = []

    async def fake_imagem(imagem, media_type, legenda, momento, hoje):
        recebidas.append((imagem, media_type))
        if imagem == b"boleto":
            return [ia.ContaExtraida(descricao="Internet", valor=99.9, vencimento="2026-10-20")]
        return []

    monkeypatch.setattr(ia, "extrair_contas_de_imagem", fake_imagem)
    resultado = _rodar(cfg, db)

    assert recebidas == [(b"boleto", "image/png"), (b"foto", "image/png")]
    assert [png for png, _ in resultado.prints] == [b"boleto"]
    assert resultado.prints[0][1].startswith("📸 Maria Eduarda")
    assert db.mensagem_processada("boleto") and db.mensagem_processada("foto")


def test_alertas_cobrem_vencendo_hoje_e_atrasadas_recentes(db):
    db.registrar_conta_a_pagar("Atrasada", date(2026, 10, 10), 5000)
    db.registrar_conta_a_pagar("Hoje", date(2026, 10, 16))
    db.registrar_conta_a_pagar("Amanha", date(2026, 10, 17))
    db.registrar_conta_a_pagar("Em tres", date(2026, 10, 19))
    db.registrar_conta_a_pagar("Em cinco", date(2026, 10, 21))
    db.registrar_conta_a_pagar("Paga", date(2026, 10, 16), 1000, pago=True)
    db.registrar_conta_a_pagar("Velha", date(2026, 8, 1), 1000)

    texto = "\n".join(contas_whatsapp.alertas_contas(db, HOJE))
    assert "ATRASADA" in texto and "Atrasada" in texto
    assert "vence HOJE" in texto and "valor a confirmar" in texto
    assert "AMANHÃ" in texto
    assert "Em tres" in texto
    assert "Em cinco" not in texto
    assert "Paga" not in texto
    assert "Velha" not in texto


def test_validacoes_de_data_e_valor():
    assert contas_whatsapp._data_valida("2026-10-15", HOJE) == date(2026, 10, 15)
    assert contas_whatsapp._data_valida("2019-01-01", HOJE) is None  # ano absurdo
    assert contas_whatsapp._data_valida("não é data", HOJE) is None
    assert contas_whatsapp._data_valida(None, HOJE) is None
    assert contas_whatsapp._centavos(183.4) == 18340
    assert contas_whatsapp._centavos(0) is None
    assert contas_whatsapp._centavos(9_999_999) is None


def test_sem_grupo_configurado_nao_tenta_ler(cfg, db):
    cfg.whatsapp_grupo = ""
    resultado = _rodar(cfg, db)
    assert resultado.status == "erro"
    assert "WHATSAPP_GRUPO" in resultado.texto


def test_planilha_mostra_as_contas_do_mes(db, cfg, tmp_path):
    from openpyxl import load_workbook

    from finbot import planilha

    db.registrar_conta_a_pagar("Luz", date(2026, 10, 15), 18340)
    db.registrar_conta_a_pagar("Internet", date(2026, 10, 20), 9990, pago=True)

    destino = planilha.gerar_planilha(db, tmp_path / "planilha.xlsx", HOJE, cfg)
    ws = load_workbook(destino)["Contas a Pagar"]
    linhas = [[c.value for c in linha] for linha in ws.iter_rows(min_row=2, max_row=3)]
    assert linhas[0][0] == "Luz" and linhas[0][3] == "Não"
    assert linhas[1][0] == "Internet" and linhas[1][3] == "Sim"
